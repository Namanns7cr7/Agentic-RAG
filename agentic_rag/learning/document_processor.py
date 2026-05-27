"""Document processor — extracts text from uploaded files.

Supported formats:
  - PDF  (.pdf)   → PyMuPDF text extraction + OCR fallback for scanned pages
  - DOCX (.docx)  → python-docx
  - TXT  (.txt)   → direct read
  - Images (.png, .jpg, .jpeg, .bmp, .tiff, .webp) → GOT-OCR2.0
  - PPTX (.pptx)  → python-pptx
  - CSV  (.csv)   → csv reader
"""

from __future__ import annotations

import csv
import io
import os
import tempfile
from pathlib import Path
from typing import Dict, List, Optional

from ..rag.utils_logger import get_logger
from .text_chunker import smart_chunk

logger = get_logger(__name__)

# Maximum file size: 50 MB
MAX_FILE_SIZE = int(os.getenv("MAX_UPLOAD_SIZE_MB", "50")) * 1024 * 1024

SUPPORTED_EXTENSIONS = {
    ".pdf", ".docx", ".txt", ".md",
    ".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp",
    ".pptx", ".csv",
}

IMAGE_EXTENSIONS = {".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".tif", ".webp"}


class DocumentProcessor:
    """Extracts text from various document types."""

    def __init__(self):
        self._ocr = None   # lazy loaded

    @property
    def ocr(self):
        if self._ocr is None:
            from .ocr_engine import get_ocr_engine
            self._ocr = get_ocr_engine()
        return self._ocr

    def process(
        self,
        file_bytes: bytes,
        filename: str,
        *,
        chunk_strategy: str = "sentence",
        chunk_size: int = 500,
    ) -> Dict:
        """Process an uploaded file and return extracted chunks.

        Returns:
            {
                "filename": str,
                "file_type": str,
                "total_chars": int,
                "num_chunks": int,
                "chunks": List[str],
                "preview": str,          # first 500 chars
            }
        """
        ext = Path(filename).suffix.lower()

        if ext not in SUPPORTED_EXTENSIONS:
            raise ValueError(
                f"Unsupported file type '{ext}'. "
                f"Supported: {', '.join(sorted(SUPPORTED_EXTENSIONS))}"
            )

        if len(file_bytes) > MAX_FILE_SIZE:
            raise ValueError(
                f"File too large ({len(file_bytes) / 1024 / 1024:.1f} MB). "
                f"Max allowed: {MAX_FILE_SIZE / 1024 / 1024:.0f} MB."
            )

        # Extract raw text
        if ext == ".pdf":
            raw_text = self._extract_pdf(file_bytes)
        elif ext == ".docx":
            raw_text = self._extract_docx(file_bytes)
        elif ext in (".txt", ".md"):
            raw_text = self._extract_text(file_bytes)
        elif ext in IMAGE_EXTENSIONS:
            raw_text = self._extract_image(file_bytes, filename)
        elif ext == ".pptx":
            raw_text = self._extract_pptx(file_bytes)
        elif ext == ".csv":
            raw_text = self._extract_csv(file_bytes)
        else:
            raise ValueError(f"No processor for '{ext}'")

        if not raw_text or not raw_text.strip():
            raise ValueError("No text could be extracted from the file.")

        # Chunk
        chunks = smart_chunk(raw_text, strategy=chunk_strategy, chunk_size=chunk_size)

        return {
            "filename": filename,
            "file_type": ext,
            "total_chars": len(raw_text),
            "num_chunks": len(chunks),
            "chunks": chunks,
            "preview": raw_text[:500],
        }

    # ── Per-format extractors ─────────────────────────────────────

    def _extract_pdf(self, data: bytes) -> str:
        """Extract text from PDF using PyMuPDF, with OCR fallback for scanned pages."""
        try:
            import fitz  # PyMuPDF
        except ImportError:
            raise RuntimeError(
                "PyMuPDF is required for PDF processing. "
                "Install it: pip install pymupdf"
            )

        doc = fitz.open(stream=data, filetype="pdf")
        pages_text: List[str] = []

        for page_num in range(len(doc)):
            page = doc.load_page(page_num)
            text = page.get_text("text").strip()

            if text and len(text) > 20:
                pages_text.append(f"--- Page {page_num + 1} ---\n{text}")
            else:
                # Page has little/no text → try OCR on rendered image
                logger.info("Page %d has no text layer — attempting OCR", page_num + 1)
                try:
                    pix = page.get_pixmap(dpi=200)
                    img_bytes = pix.tobytes("png")

                    with tempfile.NamedTemporaryFile(suffix=".png", delete=False) as tmp:
                        tmp.write(img_bytes)
                        tmp_path = tmp.name

                    ocr_text = self.ocr.extract_text(tmp_path)
                    os.unlink(tmp_path)

                    if ocr_text and "[OCR model unavailable" not in ocr_text:
                        pages_text.append(f"--- Page {page_num + 1} (OCR) ---\n{ocr_text}")
                    else:
                        pages_text.append(
                            f"--- Page {page_num + 1} ---\n"
                            f"[Scanned page — OCR model not loaded. Text could not be extracted.]"
                        )
                except Exception as exc:
                    logger.warning("OCR failed for page %d: %s", page_num + 1, exc)
                    pages_text.append(f"--- Page {page_num + 1} ---\n[OCR failed: {exc}]")

        doc.close()
        return "\n\n".join(pages_text)

    @staticmethod
    def _extract_docx(data: bytes) -> str:
        """Extract text from DOCX using python-docx."""
        try:
            from docx import Document
        except ImportError:
            raise RuntimeError(
                "python-docx is required for DOCX processing. "
                "Install it: pip install python-docx"
            )

        doc = Document(io.BytesIO(data))
        paragraphs = [p.text.strip() for p in doc.paragraphs if p.text.strip()]
        # Also extract from tables
        for table in doc.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    paragraphs.append(" | ".join(cells))
        return "\n\n".join(paragraphs)

    @staticmethod
    def _extract_text(data: bytes) -> str:
        """Extract text from plain text files."""
        for encoding in ("utf-8", "utf-16", "latin-1", "cp1252"):
            try:
                return data.decode(encoding)
            except (UnicodeDecodeError, UnicodeError):
                continue
        return data.decode("utf-8", errors="replace")

    def _extract_image(self, data: bytes, filename: str) -> str:
        """Extract text from an image using GOT-OCR2.0."""
        ext = Path(filename).suffix.lower()
        with tempfile.NamedTemporaryFile(suffix=ext, delete=False) as tmp:
            tmp.write(data)
            tmp_path = tmp.name

        try:
            text = self.ocr.extract_text(tmp_path)
            return text
        finally:
            os.unlink(tmp_path)

    @staticmethod
    def _extract_pptx(data: bytes) -> str:
        """Extract text from PPTX using python-pptx."""
        try:
            from pptx import Presentation
        except ImportError:
            raise RuntimeError(
                "python-pptx is required for PPTX processing. "
                "Install it: pip install python-pptx"
            )

        prs = Presentation(io.BytesIO(data))
        texts: List[str] = []
        for slide_num, slide in enumerate(prs.slides, 1):
            slide_texts: List[str] = []
            for shape in slide.shapes:
                if shape.has_text_frame:
                    for paragraph in shape.text_frame.paragraphs:
                        t = paragraph.text.strip()
                        if t:
                            slide_texts.append(t)
            if slide_texts:
                texts.append(f"--- Slide {slide_num} ---\n" + "\n".join(slide_texts))
        return "\n\n".join(texts)

    @staticmethod
    def _extract_csv(data: bytes) -> str:
        """Convert CSV to readable text."""
        text = data.decode("utf-8", errors="replace")
        reader = csv.reader(io.StringIO(text))
        rows: List[str] = []
        for row in reader:
            rows.append(" | ".join(cell.strip() for cell in row))
        return "\n".join(rows)


# ── Singleton ─────────────────────────────────────────────────────

_processor: Optional[DocumentProcessor] = None


def get_document_processor() -> DocumentProcessor:
    global _processor
    if _processor is None:
        _processor = DocumentProcessor()
    return _processor
